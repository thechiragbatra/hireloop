import io
import re
import uuid

from docx import Document
from pypdf import PdfReader

from ..models.schemas import Bullet, Contact, EducationItem, ExperienceItem, Profile, ProjectItem


class ResumeParser:
    def parse_pdf(self, content: bytes) -> str:
        reader = PdfReader(io.BytesIO(content))
        text = []
        for page in reader.pages:
            try:
                text.append(page.extract_text() or "")
            except Exception:
                pass
        return "\n".join(text)

    def parse_docx(self, content: bytes) -> str:
        doc = Document(io.BytesIO(content))
        text = []
        for para in doc.paragraphs:
            text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                text.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(text)

    def _extract_contact(self, text: str) -> Contact:
        email = re.findall(r"[\w\.-]+@[\w\.-]+\.\w+", text)
        phone = re.findall(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
        linkedin = re.findall(r"(https?://(?:www\.)?linkedin\.com/in/[\w\.-]+)", text)
        github = re.findall(r"(https?://(?:www\.)?github\.com/[\w\.-]+)", text)
        website = re.findall(
            r"(https?://(?!www\.linkedin|www\.github)[\w\.-]+\.[\w\.-]+/[\w\.-/]*)", text
        )
        return Contact(
            email=email[0] if email else None,
            phone=phone[0] if phone else None,
            linkedin=linkedin[0] if linkedin else None,
            github=github[0] if github else None,
            website=website[0] if website else None,
        )

    def _extract_skills(self, text: str) -> list[str]:
        common = [
            "python",
            "javascript",
            "typescript",
            "java",
            "go",
            "c++",
            "c#",
            "ruby",
            "react",
            "next.js",
            "node",
            "fastapi",
            "django",
            "flask",
            "spring",
            "sql",
            "postgres",
            "mysql",
            "mongodb",
            "redis",
            "aws",
            "gcp",
            "azure",
            "docker",
            "kubernetes",
            "terraform",
            "git",
            "linux",
            "html",
            "css",
            "tailwind",
            "pandas",
            "machine learning",
            "rest api",
            "graphql",
            "kafka",
            "pytest",
        ]
        found: list[str] = []
        tl = text.lower()
        for s in common:
            if re.search(rf"(?<!\w){re.escape(s)}(?!\w)", tl):
                found.append(s.upper() if s in {"sql", "aws", "gcp", "css", "html"} else s.title())
        return found

    @staticmethod
    def _sections(text: str) -> dict[str, list[str]]:
        aliases = {
            "experience": {
                "experience",
                "work experience",
                "professional experience",
                "employment",
            },
            "education": {"education", "academic background"},
            "projects": {"projects", "personal projects", "selected projects"},
            "skills": {"skills", "technical skills", "core competencies", "technologies"},
        }
        lookup = {name: section for section, names in aliases.items() for name in names}
        result: dict[str, list[str]] = {key: [] for key in aliases}
        current = ""
        for raw in text.splitlines():
            line = re.sub(r"\s+", " ", raw).strip(" •\t")
            if not line:
                continue
            heading = re.sub(r"[:|]+$", "", line.lower()).strip()
            if heading in lookup:
                current = lookup[heading]
                continue
            if current:
                result[current].append(line)
        return result

    def parse(self, filename: str, content: bytes) -> tuple[Profile, str]:
        ext = filename.lower().split(".")[-1]
        if ext == "pdf":
            raw_text = self.parse_pdf(content)
        elif ext == "docx":
            raw_text = self.parse_docx(content)
        else:
            raise ValueError("Unsupported file type")
        if not raw_text.strip():
            raise ValueError("No readable text found. Use a text-based PDF or DOCX file.")
        contact = self._extract_contact(raw_text)
        sections = self._sections(raw_text)
        skills = self._extract_skills(raw_text)
        known_skills = {skill.casefold() for skill in skills}
        for line in sections["skills"]:
            for skill in re.split(r"[,|•;/]+", line):
                candidate = skill.strip()
                if candidate and candidate.casefold() not in known_skills:
                    skills.append(candidate)
                    known_skills.add(candidate.casefold())

        def bullets_for(lines: list[str]) -> list[Bullet]:
            out = []
            for line in lines:
                line = re.sub(r"^[-*•▪◦]+\s*", "", line).strip()
                if line and not re.fullmatch(r"[|\-–— ]+", line):
                    out.append(Bullet(id=str(uuid.uuid4()), text=line))
            return out

        exp_bullets = bullets_for(sections["experience"])
        project_bullets = bullets_for(sections["projects"])
        education_lines = bullets_for(sections["education"])
        experience = (
            [ExperienceItem(id=str(uuid.uuid4()), company="", title="", bullets=exp_bullets)]
            if exp_bullets
            else []
        )
        projects = (
            [ProjectItem(id=str(uuid.uuid4()), name="Projects", bullets=project_bullets)]
            if project_bullets
            else []
        )
        education = []
        for item in education_lines:
            education.append(EducationItem(id=str(uuid.uuid4()), institution=item.text))
        profile = Profile(
            id=str(uuid.uuid4()),
            contact=contact,
            skills=skills,
            experience=experience,
            projects=projects,
            education=education,
        )
        return profile, raw_text
